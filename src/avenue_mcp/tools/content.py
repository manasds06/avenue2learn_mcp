"""get_course_content and read_content_file (docs/03-mcp-tools.md).

The content tree is also the RAG layer's input, so the walk here is the same
one sync_course_materials will use.
"""

from __future__ import annotations

import logging
import re
from typing import Any

from avenue_mcp.client.d2l import D2LClient
from avenue_mcp.errors import ExtractionError, InvalidRequestError, NotFoundError
from avenue_mcp.util.dates import describe, parse_d2l_date
from avenue_mcp.util.html import html_to_text

log = logging.getLogger(__name__)

# Topic types Brightspace reports. Only File topics have a body to download;
# Link and Page topics do not, and asking for one returns an error.
TOPIC_FILE = 1
TOPIC_LINK = 3

_TEXT_MIMES = ("text/", "application/json", "application/xml")


def _is_downloadable(topic: dict[str, Any]) -> bool:
    """Whether read_content_file can do anything with this topic.

    Surfaced as a flag so the model doesn't have to guess and burn a tool call
    discovering that a link topic has no file.
    """
    if topic.get("TypeIdentifier") == "Link":
        return False
    url = topic.get("Url") or ""
    if not url:
        return False
    # An external link topic points off-instance; there's nothing to fetch
    # from the content API for it.
    return not url.startswith(("http://", "https://"))


def _normalize_topic(topic: dict[str, Any], tz: str) -> dict[str, Any]:
    url = topic.get("Url") or ""
    file_name = url.rsplit("/", 1)[-1] if url else ""
    return {
        "id": topic.get("Id"),
        "title": topic.get("Title") or "",
        "type": topic.get("TypeIdentifier") or "Topic",
        "file_name": file_name,
        "url": url,
        "last_modified": describe(parse_d2l_date(topic.get("LastModifiedDate")), tz),
        "is_downloadable": _is_downloadable(topic),
    }


async def _walk_module(
    client: D2LClient,
    org_unit_id: int,
    module: dict[str, Any],
    *,
    depth: int,
    max_depth: int,
    visited: set[int],
    counter: dict[str, int],
    tz: str,
) -> dict[str, Any]:
    """Recursively expand a module.

    Depth-capped and cycle-guarded: recursive tree-walking against a remote
    API is a good way to write an accidental infinite loop if the data ever
    contains a cycle.
    """
    module_id = module.get("Id")
    node: dict[str, Any] = {
        "id": module_id,
        "title": module.get("Title") or "",
        "topics": [],
        "modules": [],
    }

    if module_id is None or module_id in visited or depth >= max_depth:
        if depth >= max_depth:
            node["truncated"] = "max_depth reached"
        return node
    visited.add(module_id)

    try:
        children = await client.get_json(
            "le", f"{org_unit_id}/content/modules/{module_id}/structure/"
        )
    except NotFoundError:
        return node

    for child in children if isinstance(children, list) else []:
        if not isinstance(child, dict):
            continue
        if child.get("Type") == 0 or "Structure" in child or child.get("TypeIdentifier") == "Module":
            node["modules"].append(
                await _walk_module(
                    client, org_unit_id, child,
                    depth=depth + 1, max_depth=max_depth,
                    visited=visited, counter=counter, tz=tz,
                )
            )
        else:
            node["topics"].append(_normalize_topic(child, tz))
            counter["topics"] += 1

    return node


async def get_course_content(
    client: D2LClient,
    *,
    org_unit_id: int,
    max_depth: int = 10,
) -> dict[str, Any]:
    """The course's Content tree — structure only, no file contents."""
    tz = client.settings.timezone
    roots = await client.get_json("le", f"{org_unit_id}/content/root/")

    counter = {"topics": 0}
    visited: set[int] = set()
    modules: list[dict[str, Any]] = []

    for root in roots if isinstance(roots, list) else []:
        if not isinstance(root, dict):
            continue
        modules.append(
            await _walk_module(
                client, org_unit_id, root,
                depth=0, max_depth=max_depth,
                visited=visited, counter=counter, tz=tz,
            )
        )

    return {
        "org_unit_id": org_unit_id,
        "modules": modules,
        "topic_count": counter["topics"],
        "module_count": len(visited),
    }


def _extract_text(data: bytes, mime: str, file_name: str) -> tuple[str, int | None]:
    """Extract text and a page/slide count from file bytes.

    PDF/DOCX/PPTX need the optional [rag] extras. When they're missing we say
    so with an install hint rather than failing opaquely — the core server is
    installable without onnxruntime and friends by design.
    """
    lowered = file_name.lower()

    if lowered.endswith(".pdf") or "pdf" in mime:
        try:
            import pymupdf
        except ImportError as exc:
            raise ExtractionError(
                "Reading PDFs needs the optional extras. Install with: pip install -e \".[rag]\""
            ) from exc
        with pymupdf.open(stream=data, filetype="pdf") as doc:
            # Page markers so the model can cite precisely.
            pages = [f"[page {i + 1}]\n{page.get_text()}" for i, page in enumerate(doc)]
            return "\n\n".join(pages), len(pages)

    if lowered.endswith(".pptx"):
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise ExtractionError(
                "Reading PPTX needs the optional extras. Install with: pip install -e \".[rag]\""
            ) from exc
        import io

        deck = Presentation(io.BytesIO(data))
        slides = []
        for i, slide in enumerate(deck.slides):
            parts = [sh.text for sh in slide.shapes if getattr(sh, "has_text_frame", False)]
            # Speaker notes often carry the actual explanation the slide omits.
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    parts.append(f"[notes] {notes}")
            slides.append(f"[slide {i + 1}]\n" + "\n".join(p for p in parts if p))
        return "\n\n".join(slides), len(slides)

    if lowered.endswith(".docx"):
        try:
            import docx
        except ImportError as exc:
            raise ExtractionError(
                "Reading DOCX needs the optional extras. Install with: pip install -e \".[rag]\""
            ) from exc
        import io

        document = docx.Document(io.BytesIO(data))
        return "\n".join(p.text for p in document.paragraphs if p.text.strip()), None

    if lowered.endswith((".html", ".htm")) or "html" in mime:
        return html_to_text(data.decode("utf-8", errors="replace")), None

    if lowered.endswith((".txt", ".md", ".csv")) or any(m in mime for m in _TEXT_MIMES):
        return data.decode("utf-8", errors="replace"), None

    raise ExtractionError(
        f"Can't extract text from {file_name or 'this file'} (type: {mime or 'unknown'}). "
        f"Supported: PDF, DOCX, PPTX, HTML, TXT."
    )


async def read_content_file(
    client: D2LClient,
    *,
    org_unit_id: int,
    topic_id: int,
    max_chars: int = 50_000,
) -> dict[str, Any]:
    """Download one Content file and return its extracted text."""
    topic = await client.get_json("le", f"{org_unit_id}/content/topics/{topic_id}")
    if not isinstance(topic, dict):
        raise NotFoundError(f"Topic {topic_id} not found in course {org_unit_id}.")

    if not _is_downloadable(topic):
        raise InvalidRequestError(
            f"Topic {topic_id} ({topic.get('Title') or 'untitled'}) is a link or page, "
            f"not a file — there's nothing to download. Use get_course_content to find "
            f"topics with is_downloadable: true."
        )

    headers, payload = await client.download_file(
        "le", f"{org_unit_id}/content/topics/{topic_id}/file"
    )
    data = payload if isinstance(payload, bytes) else payload.read_bytes()

    url = topic.get("Url") or ""
    file_name = url.rsplit("/", 1)[-1] or f"topic-{topic_id}"
    disposition = headers.get("content-disposition", "")
    if match := re.search(r'filename\*?=(?:UTF-8\'\')?"?([^";]+)', disposition):
        file_name = match.group(1)
    mime = headers.get("content-type", "").split(";")[0].strip()

    text, page_count = _extract_text(data, mime, file_name)

    # A multi-page document that yields almost nothing is a scanned image.
    # Say so rather than returning an empty string that reads as "no content".
    quality = "ok"
    if page_count and page_count > 1 and len(text.strip()) < 100:
        quality = "poor"

    truncated = len(text) > max_chars
    return {
        "topic_id": topic_id,
        "title": topic.get("Title") or "",
        "file_name": file_name,
        "mime_type": mime,
        "page_count": page_count,
        "text": text[:max_chars],
        "truncated": truncated,
        "extraction_quality": quality,
        "note": (
            "This document appears to be scanned images — almost no text could be "
            "extracted. OCR is not supported."
            if quality == "poor"
            else None
        ),
    }
